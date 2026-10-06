"""File report (TASKS B27): TXT/DOCX/PDF, size limit, paragraph chunks, aggregation."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from app.api import file_report as fr
from app.models.result import CheckResult, ClaimResult
from app.pipeline import orchestrator


def test_chunking_keeps_paragraphs_under_limit() -> None:
    text = "\n".join(f"paragraph {i} " + "x" * 300 for i in range(40))
    chunks = fr.chunk_paragraphs(text, 4000)
    assert all(len(c) <= 4000 for c in chunks) and len(chunks) > 1
    assert "\n".join(chunks).count("paragraph") == 40


def test_docx_text_extracted() -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("first line")
    d.add_paragraph("second line")
    buf = io.BytesIO()
    d.save(buf)
    assert fr.extract_text("a.docx", buf.getvalue()).split("\n") == ["first line", "second line"]


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch):
    from app.api import check as check_api
    from app.main import app

    async def fake_run(text, channel="web", lang_hint=None):
        return CheckResult(check_id="e" * 32, status="ok", lang="en", disclaimer="d", expires_at="2026-10-07T00:00:00Z",
                           claims=[ClaimResult(index=0, type="hadith", span="s", span_start=0, span_end=1, lang="en",
                                               verdict="not_established")])

    monkeypatch.setattr(orchestrator, "run_check", fake_run)
    check_api.limiter._hits.clear()
    with TestClient(app) as c:
        yield c


def test_txt_upload_aggregates(client: TestClient) -> None:
    r = client.post("/api/v1/check/file", files={"file": ("m.txt", "line one\nline two".encode(), "text/plain")})
    body = r.json()
    assert r.status_code == 200 and body["summary"] == {"not_established": 1} and body["checks"][0]["check_id"]


def test_too_large_is_413(client: TestClient) -> None:
    big = b"x" * (5 * 1024 * 1024 + 1)
    assert client.post("/api/v1/check/file", files={"file": ("m.txt", big, "text/plain")}).status_code == 413


def test_unsupported_type_is_415(client: TestClient) -> None:
    assert client.post("/api/v1/check/file", files={"file": ("m.exe", b"MZ", "application/octet-stream")}).status_code == 415
